"""Собрать эталон для золотых тестов из выпущенных протоколов `.xlsm`.

    python tools/extract_golden.py ~/protocols/2026-08

Читает посчитанные Excel значения (data_only=True) и раскладку определяет по
формулам (data_only=False): короткий шаблон — 3 строки измерений, полный — 9.
Длительность пролива берётся прямо из формулы `ROUNDUP(ACnn*<t>/3600;3)`,
поэтому менять её здесь руками не нужно.

Результат кладётся в backend/apps/verification/tests/golden/.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).parent.parent / "backend/apps/verification/tests/golden/water_meter_protocols.json"
COLUMNS = {
    "q": "AC", "reading_start": "AJ", "reading_end": "AQ",
    "v_standard": "AW", "v_meter": "BD", "error_pct": "BK",
}


def extract(path: Path) -> dict:
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    protocol, data = values["Протокол"], values["Данные"]
    protocol_formulas = formulas["Протокол"]

    last_row = 61 if protocol_formulas["AC61"].value else 55
    measurements = []
    for row in range(53, last_row + 1):
        formula = str(protocol_formulas[f"AW{row}"].value or "")
        seconds = re.search(r"AC\d+\*(\d+)/3600", formula)
        measurements.append(
            {"row": row, "seconds": int(seconds.group(1)) if seconds else None}
            | {key: protocol[f"{column}{row}"].value for key, column in COLUMNS.items()}
        )

    return {
        "file": path.name,
        "protocol": protocol["A5"].value,
        "meter_type": data["B1"].value,
        "meter_class": protocol["V53"].value,
        "layout": "extended9" if last_row == 61 else "compact3",
        "measurements": measurements,
    }


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 2

    source = Path(argv[1]).expanduser()
    files = sorted(source.glob("*.xlsm")) if source.is_dir() else [source]
    if not files:
        print(f"Не нашёл .xlsm в {source}")
        return 1

    extracted = []
    for path in files:
        try:
            extracted.append(extract(path))
        except Exception as exc:  # noqa: BLE001 — отчёт важнее падения
            print(f"  пропущен {path.name}: {exc}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(extracted, ensure_ascii=False, indent=1), encoding="utf-8")

    rows = sum(len(item["measurements"]) for item in extracted)
    layouts = {item["layout"] for item in extracted}
    print(f"Протоколов: {len(extracted)}, строк измерений: {rows}, раскладки: {sorted(layouts)}")
    print(f"Записано: {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))

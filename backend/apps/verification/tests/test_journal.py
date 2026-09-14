"""Веб-журнал, выгрузка в формат ФИФ и экран нормоконтроля."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.models import User
from apps.verification import journal, numbering
from apps.verification import measurements as service
from apps.verification.calculators import water_meter as wm
from apps.verification.models import Client, ProtocolStatus, Site, VerificationStatus
from apps.verification.tests import factories as f

LIMITS = {
    "q_min": "0.03", "q_transition_a": "0.15", "q_transition_b": "0.12",
    "q_nominal": "1.5", "q_max": "3",
    "error_below_transition": "5", "error_above_transition": "2",
}
ROWS = [
    {"flow_rate": "0.030", "reading_start": "1091.763", "reading_end": "1091.769",
     "volume_standard": "0.0063"},
    {"flow_rate": "0.132", "reading_start": "1091.775", "reading_end": "1091.789",
     "volume_standard": "0.0142"},
    {"flow_rate": "0.666", "reading_start": "1091.807", "reading_end": "1091.830",
     "volume_standard": "0.0226"},
]


class JournalTestCaseMixin:
    def prepare(self):
        self.family = f.make_family()
        self.employee = f.make_employee()
        f.attest(self.employee, self.family)
        self.user = User.objects.create_user("j@ei.test", "pw")
        self.user.employee = self.employee
        self.employee.user = self.user
        self.employee.save()

        self.client_obj = Client.objects.create(name="Частное лицо")
        self.site = Site.objects.create(address="г.Киров, ул.Ленинградская 1а-70")
        self.si_type = f.make_si_type(
            self.family, registry_number="16078-13", name="СХВ-15", limits=LIMITS
        )

    def make_verification(self, *, day: int, serial: str, measure: bool = True):
        verification = f.make_verification(
            self.family, self.employee, day=day, serial=serial, si_type=self.si_type
        )
        verification.instrument.owner = self.client_obj
        verification.instrument.site = self.site
        verification.instrument.manufacture_year = 2014
        verification.instrument.save()
        verification.temperature = Decimal("21")
        verification.humidity = Decimal("53.1")
        verification.pressure = Decimal("100")
        verification.save()

        if measure:
            service.apply(
                verification,
                {"layout": wm.LAYOUT_COMPACT, "meter_class": wm.CLASS_B,
                 "unit_type": "х/в", "rows": ROWS},
            )
            verification.refresh_from_db()
        return verification


class JournalApiTestCase(JournalTestCaseMixin, TestCase):
    def setUp(self) -> None:
        self.prepare()
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.first = self.make_verification(day=3, serial="SN003")
        self.second = self.make_verification(day=7, serial="SN007")

    def rows(self, **params):
        response = self.api.get(reverse("journal"), params)
        assert response.status_code == 200, response.data
        return response.data["results"]

    def test_journal_lists_newest_first(self):
        rows = self.rows()
        assert [row["serial_number"] for row in rows] == ["SN007", "SN003"]

    def test_filter_by_date_range(self):
        rows = self.rows(date_from="2026-03-05", date_to="2026-03-10")
        assert [row["serial_number"] for row in rows] == ["SN007"]

    def test_search_by_serial_and_by_address(self):
        assert len(self.rows(q="SN003")) == 1
        assert len(self.rows(q="Ленинградская")) == 2
        assert len(self.rows(q="несуществующее")) == 0

    def test_filter_by_suitability(self):
        rows = [dict(row) for row in ROWS]
        rows[2]["volume_standard"] = "0.0280"
        service.apply(
            self.second,
            {"layout": wm.LAYOUT_COMPACT, "meter_class": wm.CLASS_B, "rows": rows},
        )
        assert [row["serial_number"] for row in self.rows(suitable="false")] == ["SN007"]
        assert [row["serial_number"] for row in self.rows(suitable="true")] == ["SN003"]

    def test_row_carries_what_the_table_shows(self):
        row = self.rows(q="SN007")[0]

        assert row["si_name"] == "СХВ-15"
        assert row["address"] == "г.Киров, ул.Ленинградская 1а-70"
        assert row["verifier"] == "Стариков А. А."
        assert row["journal_note"].startswith("Поверен в диапазоне расхода")
        assert row["needs_review"] is False

    def test_scanned_rows_awaiting_review_can_be_filtered(self):
        scanned = [dict(row) | {"confidence": 0.5} for row in ROWS]
        service.apply(
            self.second,
            {"layout": wm.LAYOUT_COMPACT, "source": service.SOURCE_SCAN, "rows": scanned},
        )
        assert [row["serial_number"] for row in self.rows(needs_review="true")] == ["SN007"]


class JournalExportTestCase(JournalTestCaseMixin, TestCase):
    def setUp(self) -> None:
        self.prepare()
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.verification = self.make_verification(day=7, serial="12470437")
        protocol = f.make_protocol(self.verification)
        numbering.assign_numbers(numbering.scope_for(self.verification))
        protocol.refresh_from_db()
        self.protocol = protocol

    def sheet(self):
        response = self.api.get(reverse("journal_export"))
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        return workbook[journal.SHEET_NAME]

    def test_export_has_the_original_49_columns_and_headers(self):
        sheet = self.sheet()
        assert sheet.max_column == 49
        assert sheet.cell(row=1, column=1).value == "Номер протокола"
        assert sheet.cell(row=1, column=3).value.startswith("Регистрационный номер типа СИ")

    def test_values_land_in_the_fif_columns(self):
        sheet = self.sheet()
        row = 2

        assert sheet.cell(row=row, column=1).value == self.protocol.full_number
        assert sheet.cell(row=row, column=2).value == "1"
        assert sheet.cell(row=row, column=3).value == "16078-13"
        assert sheet.cell(row=row, column=6).value == "СХВ-15"
        assert sheet.cell(row=row, column=8).value == "12470437"
        assert sheet.cell(row=row, column=14).value == "Пригодно"
        assert sheet.cell(row=row, column=15).value == "21.0 °C"   # как в прежнем журнале: «23.4 °C»

    def test_address_goes_to_the_owner_column_as_before(self):
        """В прежнем журнале адрес поверки лежал в графе «Владелец СИ»."""
        sheet = self.sheet()
        assert sheet.cell(row=2, column=33).value == "г.Киров, ул.Ленинградская 1а-70"
        assert sheet.cell(row=2, column=48).value == "Частное лицо"

    def test_other_info_carries_the_flow_range_note(self):
        sheet = self.sheet()
        assert sheet.cell(row=2, column=42).value.startswith("Поверен в диапазоне расхода")

    def test_unit_type_lands_in_the_type_column(self):
        sheet = self.sheet()
        assert sheet.cell(row=2, column=49).value == "х/в"

    def test_export_is_ordered_chronologically(self):
        """Журнал читается как хронология, поэтому выгрузка по возрастанию даты."""
        self.make_verification(day=3, serial="SN003")
        sheet = self.sheet()
        dates = [sheet.cell(row=r, column=10).value for r in (2, 3)]
        assert dates[0] < dates[1]


class NormocontrolTestCase(JournalTestCaseMixin, TestCase):
    def setUp(self) -> None:
        self.prepare()
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.third = self.make_verification(day=3, serial="SN003")
        self.seventh = self.make_verification(day=7, serial="SN007")
        for verification in (self.seventh, self.third):
            f.make_protocol(verification)
        self.scope = numbering.scope_for(self.third)

    def test_scopes_show_what_is_waiting(self):
        response = self.api.get(reverse("normocontrol_scopes"))
        assert response.status_code == 200
        scope = response.data[0]

        assert scope["drafts"] == 2
        assert scope["sealed_high_water"] == 0
        assert scope["chronology_breaks"] == []

    def test_preview_changes_nothing(self):
        response = self.api.post(reverse("numbering_preview", args=[self.scope.pk]))

        assert response.status_code == 200
        assert response.data["changed"] is True
        assert len(response.data["assigned"]) == 2
        self.third.protocol.refresh_from_db()
        assert self.third.protocol.seq is None

    def test_apply_assigns_numbers_in_chronological_order(self):
        response = self.api.post(reverse("numbering_apply", args=[self.scope.pk]))

        assert response.status_code == 200
        assert response.data["assigned"] == 2
        self.third.protocol.refresh_from_db()
        self.seventh.protocol.refresh_from_db()
        assert self.third.protocol.seq == 1
        assert self.seventh.protocol.seq == 2

    def test_signing_a_batch_reports_what_went_through(self):
        self.api.post(reverse("numbering_apply", args=[self.scope.pk]))
        self.third.protocol.refresh_from_db()
        self.seventh.protocol.refresh_from_db()

        response = self.api.post(
            reverse("sign_protocols"),
            {"protocols": [self.third.protocol.pk, self.seventh.protocol.pk]},
            format="json",
        )

        assert response.status_code == 200
        assert len(response.data["signed"]) == 2
        assert response.data["refused"] == []
        self.third.protocol.refresh_from_db()
        assert self.third.protocol.status == ProtocolStatus.SIGNED

    def test_one_bad_protocol_does_not_sink_the_batch(self):
        """Без номера подписать нельзя, но остальные пройти должны."""
        self.api.post(reverse("numbering_apply", args=[self.scope.pk]))
        self.third.protocol.refresh_from_db()
        orphan = f.make_protocol(self.make_verification(day=9, serial="SN009"))

        response = self.api.post(
            reverse("sign_protocols"),
            {"protocols": [self.third.protocol.pk, orphan.pk]},
            format="json",
        )

        assert len(response.data["signed"]) == 1
        assert len(response.data["refused"]) == 1
        assert "без номера" in response.data["refused"][0]["detail"]

    def test_signing_without_an_employee_card_is_refused_clearly(self):
        outsider = User.objects.create_user("outsider@ei.test", "pw")
        api = APIClient()
        api.force_authenticate(outsider)

        response = api.post(
            reverse("sign_protocols"), {"protocols": [self.third.protocol.pk]}, format="json"
        )

        assert response.status_code == 400
        assert "карточки сотрудника" in response.data["detail"]

    def test_status_of_verification_is_visible_in_the_journal(self):
        assert self.third.status == VerificationStatus.READY

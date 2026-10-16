"""Текст из файлов: PDF, Word, Excel, CSV/TXT/MD, HTML — и нарезка на куски для поиска.

Сканы без текстового слоя (PDF-картинка) текста не дают — это честно
сообщается в ``extract_error``: такой документ нужно либо загрузить
картинкой в чат (модель с vision прочтёт), либо распознать заранее.
"""

from __future__ import annotations

import csv
import html
import io
import mimetypes
import re
from html.parser import HTMLParser
from pathlib import Path

MAX_TEXT = 3_000_000  # больше — режем: такой файл всё равно читают поиском


class ExtractError(Exception):
    pass


def guess_mime(name: str, declared: str = "") -> str:
    if declared and declared != "application/octet-stream":
        return declared
    return mimetypes.guess_type(name)[0] or "application/octet-stream"


def extract_text(data: bytes, name: str, mime: str = "") -> str:
    """Вернуть текст документа или поднять ExtractError с понятной причиной."""
    suffix = Path(name).suffix.lower()
    mime = guess_mime(name, mime)

    if suffix == ".pdf" or mime == "application/pdf":
        text = _pdf(data)
    elif suffix == ".docx" or "wordprocessingml" in mime:
        text = _docx(data)
    elif suffix in (".xlsx", ".xlsm") or "spreadsheetml" in mime:
        text = _xlsx(data)
    elif suffix == ".csv":
        text = _csv(data)
    elif suffix in (".html", ".htm") or mime == "text/html":
        text = html_to_text(_decode(data))
    elif suffix in (".doc", ".xls", ".rtf"):
        raise ExtractError(
            f"Старый формат {suffix} не читается — сохраните файл как "
            f"{'.docx' if suffix != '.xls' else '.xlsx'} и загрузите снова"
        )
    elif mime.startswith("image/"):
        return ""
    elif mime.startswith("text/") or suffix in (".txt", ".md", ".json", ".xml", ".yaml", ".yml", ".typ"):
        text = _decode(data)
    else:
        raise ExtractError(f"Не умею читать файлы типа {suffix or mime}")

    text = normalize(text)
    if not text.strip():
        raise ExtractError(
            "В файле нет текстового слоя (похоже на скан). Загрузите его в чат "
            "картинками — модель с распознаванием изображений прочтёт"
        )
    return text[:MAX_TEXT]


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("pypdf не установлен") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            pages.append(f"[стр. {number}]\n{page.extract_text() or ''}")
        return "\n\n".join(pages) if any(p.split("\n", 1)[1].strip() for p in pages) else ""
    except Exception as exc:  # noqa: BLE001 — битый PDF не должен ронять загрузку
        raise ExtractError(f"PDF не читается: {exc}") from exc


def _docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("python-docx не установлен") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"Word-файл не читается: {exc}") from exc

    # Абзацы и таблицы — в порядке следования в документе.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if style.startswith("heading") or style.startswith("заголовок"):
                text = "## " + text
            lines.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if not cells or cells[-1] != value:  # объединённые ячейки повторяются
                        cells.append(value)
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"Excel-файл не читается: {exc}") from exc
    parts: list[str] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            while values and values[-1] == "":
                values.pop()
            if values:
                rows.append("\t".join(values))
            if len(rows) >= 20000:
                rows.append("… (дальше строки обрезаны)")
                break
        if rows:
            parts.append(f"## Лист «{sheet.title}»\n" + "\n".join(rows))
    workbook.close()
    return "\n\n".join(parts)


def _csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text), dialect))


class _TextHTML(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")
        if tag in ("h1", "h2", "h3"):
            self.parts.append("## ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_to_text(source: str) -> str:
    parser = _TextHTML()
    try:
        parser.feed(source)
    except Exception:  # noqa: BLE001 — кривой HTML: берём что успели
        pass
    text = "".join(parser.parts)
    text = re.sub(r"[ \t]+", " ", html.unescape(text))
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# Нарезка для поиска
# ---------------------------------------------------------------------------
CHUNK_SIZE = 1800
CHUNK_OVERLAP = 200
HEADING_RE = re.compile(r"^(#{1,4} .+|\d+(\.\d+)*\.?\s+[А-ЯA-Z].{2,120}|\[стр\. \d+\])$")


def chunk_text(text: str, size: int = CHUNK_SIZE) -> list[tuple[str, str]]:
    """Нарезать на куски ~size символов по границам абзацев. Возвращает (заголовок, текст).

    Заголовок куска — последний заголовок раздела перед его первым абзацем:
    по нему агент понимает, из какого пункта документа фрагмент.
    """
    pieces: list[tuple[str, str]] = []
    heading = ""
    for paragraph in text.split("\n"):
        stripped = paragraph.strip()
        if HEADING_RE.match(stripped):
            heading = stripped.lstrip("# ")[:300]
        while len(paragraph) > size:  # очень длинная строка (таблица Excel без переносов)
            pieces.append((heading, paragraph[:size]))
            paragraph = paragraph[size:]
        pieces.append((heading, paragraph))

    chunks: list[tuple[str, str]] = []
    buffer: list[str] = []
    length = 0
    current = ""
    fresh = False  # есть ли в буфере что-то кроме перекрытия с прошлым куском
    for piece_heading, paragraph in pieces:
        if fresh and length + len(paragraph) > size:
            body = "\n".join(buffer).strip()
            if body:
                chunks.append((current, body))
            tail = body[-CHUNK_OVERLAP:] if len(body) > CHUNK_OVERLAP else ""
            buffer, length, fresh = ([tail] if tail else []), len(tail), False
        if not fresh and paragraph.strip():
            current = piece_heading
        buffer.append(paragraph)
        length += len(paragraph) + 1
        fresh = fresh or bool(paragraph.strip())
    body = "\n".join(buffer).strip()
    if fresh and body:
        chunks.append((current, body))
    return chunks

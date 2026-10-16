"""Ассистент: агентный цикл, инструменты, подтверждение действий, контекст, MCP."""

from __future__ import annotations

import base64
import io
import json
import shutil
from unittest import mock

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework.test import APIClient

from apps.assistant import agent, extract, knowledge, llm, mcp_client, web
from apps.assistant.models import (
    Attachment,
    Conversation,
    KnowledgeFile,
    MemoryItem,
    Message,
    PendingAction,
    Run,
    Scope,
)
from apps.assistant.prompt import build_system_prompt
from apps.catalog.models import MeasurementFamily, SiType
from apps.core.models import Role, User

LIMITS = {
    "q_min": 0.03, "q_transition_a": 0.15, "q_transition_b": 0.12, "q_nominal": 1.5,
    "q_max": 3, "error_below_transition": 5, "error_above_transition": 2,
}


class FakeLLM:
    """Сценарий ответов модели: список Completion по порядку вызовов."""

    def __init__(self, *steps: llm.Completion) -> None:
        self.steps = list(steps)
        self.calls: list[dict] = []

    def __call__(self, model, messages, *, tools=None, on_text=None, should_stop=None, **kw):
        self.calls.append({"model": model, "messages": messages, "tools": tools})
        step = self.steps.pop(0)
        if on_text and step.text:
            on_text(step.text)
        return step


def call(name: str, args: dict, call_id: str = "c1") -> llm.Completion:
    return llm.Completion(tool_calls=[llm.ToolCall(id=call_id, name=name, arguments=json.dumps(args, ensure_ascii=False))])


def answer(text: str) -> llm.Completion:
    return llm.Completion(text=text, usage={"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15, "cost": 0.001})


class Base(TestCase):
    def setUp(self) -> None:
        self.family = MeasurementFamily.objects.create(code="water-meter", name="Счётчики воды", calculator_key="water_meter")
        self.verifier = User.objects.create_user("v@ei.test", "pw", role=Role.VERIFIER)
        self.metrologist = User.objects.create_user("m@ei.test", "pw", role=Role.METROLOGIST)
        self.conversation = Conversation.objects.create(user=self.verifier, model="openrouter:test/model")
        # MCP-серверов нет — список инструментов строится без сети
        patcher = mock.patch("apps.assistant.llm.available_models", return_value=[])
        patcher.start()
        self.addCleanup(patcher.stop)

    def client_for(self, user) -> APIClient:
        client = APIClient()
        client.force_authenticate(user)
        return client

    def send(self, text: str, fake: FakeLLM, user=None):
        with mock.patch("apps.assistant.llm.chat", fake):
            return self.client_for(user or self.verifier).post(
                reverse("assistant_send", args=[self.conversation.pk]), {"content": text}, format="json",
            )


# ---------------------------------------------------------------------------
class AgentLoopTests(Base):
    def test_tool_then_answer(self):
        kf = KnowledgeFile.objects.create(title="Регламент поверки манометров", category="rules", text="")
        kf.text = "Манометры поверяем по МИ 2124-90. Межповерочный интервал 12 месяцев."
        knowledge.index_file(kf)
        fake = FakeLLM(call("context_search", {"query": "манометр"}), answer("Поверяем по МИ 2124-90."))
        response = self.send("По какой методике поверяем манометры?", fake)
        assert response.status_code == 201, response.data
        run = Run.objects.get()
        assert run.status == Run.Status.DONE, run.error
        roles = list(self.conversation.messages.values_list("role", flat=True))
        assert roles == ["user", "assistant", "tool", "assistant"]
        tool_msg = self.conversation.messages.get(role="tool")
        assert "МИ 2124-90" in tool_msg.content
        assert tool_msg.ui["sources"][0]["title"] == "Регламент поверки манометров"
        # Во второй вызов модель получила результат инструмента
        second = fake.calls[1]["messages"]
        assert second[-1]["role"] == "tool" and second[-2]["tool_calls"][0]["function"]["name"] == "context_search"
        assert run.usage["total_tokens"] == 15
        self.conversation.refresh_from_db()
        assert self.conversation.title.startswith("По какой методике")

    def test_tool_error_is_shown_to_model(self):
        fake = FakeLLM(call("context_read", {"file_id": 999}), answer("Файла нет."))
        self.send("прочитай файл 999", fake)
        tool_msg = self.conversation.messages.get(role="tool")
        assert tool_msg.ui["error"] is True
        assert tool_msg.content.startswith("Ошибка:")
        assert Run.objects.get().status == Run.Status.DONE

    def test_unknown_tool_and_bad_json(self):
        bad = llm.Completion(tool_calls=[llm.ToolCall(id="x", name="context_search", arguments="{не json")])
        fake = FakeLLM(bad, call("nope", {}, "y"), answer("ок"))
        self.send("тест", fake)
        contents = list(self.conversation.messages.filter(role="tool").values_list("content", flat=True))
        assert "не JSON" in contents[0] and "Нет инструмента" in contents[1]

    def test_step_limit_forces_summary(self):
        from apps.assistant.models import AssistantSettings

        settings_obj = AssistantSettings.current()
        settings_obj.max_steps = 2
        settings_obj.save()
        fake = FakeLLM(
            call("ei_reference", {"what": "families"}, "a"), call("ei_reference", {"what": "families"}, "b"),
            answer("Итог."),
        )
        self.send("ищи бесконечно", fake)
        assert fake.calls[-1]["tools"] is None
        assert "Лимит шагов" in fake.calls[-1]["messages"][-1]["content"]

    def test_llm_error_marks_run(self):
        def boom(*a, **kw):
            raise llm.LLMError("openrouter ответил 401: bad key")

        with mock.patch("apps.assistant.llm.chat", boom):
            self.client_for(self.verifier).post(
                reverse("assistant_send", args=[self.conversation.pk]), {"content": "привет"}, format="json",
            )
        run = Run.objects.get()
        assert run.status == Run.Status.ERROR and "401" in run.error

    def test_cannot_send_while_running(self):
        Run.objects.create(conversation=self.conversation, model="m", status=Run.Status.RUNNING)
        response = self.send("ещё", FakeLLM())
        assert response.status_code == 400

    def test_run_state_and_cancel(self):
        run = Run.objects.create(conversation=self.conversation, model="m", status=Run.Status.RUNNING)
        client = self.client_for(self.verifier)
        state = client.get(reverse("assistant_run", args=[run.pk]) + "?after=0").data
        assert state["run"]["active"] is True
        client.post(reverse("assistant_run_cancel", args=[run.pk]))
        run.refresh_from_db()
        assert run.status == Run.Status.CANCELLING
        # чужой прогон не виден
        assert self.client_for(self.metrologist).get(reverse("assistant_run", args=[run.pk])).status_code == 404


class BuildMessagesTests(Base):
    def test_orphan_tool_calls_and_deferred_events(self):
        c = self.conversation
        Message.objects.create(conversation=c, role="user", content="привет")
        Message.objects.create(conversation=c, role="assistant", content="", tool_calls=[
            {"id": "a", "type": "function", "function": {"name": "x", "arguments": "{}"}},
            {"id": "b", "type": "function", "function": {"name": "x", "arguments": "{}"}},
        ])
        Message.objects.create(conversation=c, role="tool", tool_call_id="a", content="A")
        Message.objects.create(conversation=c, role="event", content="[Событие] подтверждено")
        # результата для b нет — прогон прерван
        Message.objects.create(conversation=c, role="user", content="дальше")
        out = agent.build_messages(c, vision=False)
        assert [m["role"] for m in out] == ["user", "assistant", "tool", "tool", "user", "user"]
        assert out[3]["tool_call_id"] == "b" and "прерван" in out[3]["content"]
        assert out[4]["content"].startswith("[Событие]")

    def test_image_only_for_vision_models(self):
        from PIL import Image

        buffer = io.BytesIO()
        Image.new("RGB", (20, 20), "red").save(buffer, format="PNG")
        attachment = Attachment(user=self.verifier, name="шильдик.png", mime="image/png", size=1)
        attachment.file.save("x.png", SimpleUploadedFile("x.png", buffer.getvalue()), save=False)
        attachment.save()
        message = Message.objects.create(conversation=self.conversation, role="user", content="что на фото?")
        message.attachments.add(attachment)
        with_vision = agent.build_messages(self.conversation, vision=True)[0]["content"]
        assert with_vision[1]["type"] == "image_url"
        assert with_vision[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
        without = agent.build_messages(self.conversation, vision=False)[0]["content"]
        assert isinstance(without, str) and "не видит изображения" in without


# ---------------------------------------------------------------------------
class ActionTests(Base):
    def propose(self, **overrides):
        args = {
            "registry_number": "77777-21", "name": "Счётчик ТЕСТ-15", "manufacturer": "ООО Завод",
            "verification_interval_months": 72, "family": self.family.pk, "limits": LIMITS,
            "sources": ["ФГИС: get_type 77777-21", "Описание типа, табл. 2"],
        }
        args.update(overrides)
        fake = FakeLLM(call("ei_si_type_propose", args), answer("Готово, подтвердите."))
        self.send("заведи тип", fake)
        return self.conversation.messages.get(role="tool")

    def test_propose_creates_pending_action_not_type(self):
        tool_msg = self.propose()
        assert not SiType.objects.filter(registry_number="77777-21").exists()
        action = PendingAction.objects.get()
        assert tool_msg.ui["action"] == action.pk
        detail = self.client_for(self.verifier).get(reverse("assistant_conversation", args=[self.conversation.pk])).data
        tool_json = [m for m in detail["messages"] if m["role"] == "tool"][0]
        assert tool_json["ui"]["action"]["status"] == "pending"

    def test_confirm_creates_type_waiting_for_review(self):
        self.propose()
        action = PendingAction.objects.get()
        response = self.client_for(self.verifier).post(reverse("assistant_action_decide", args=[action.pk, "confirm"]))
        assert response.status_code == 200, response.data
        si_type = SiType.objects.get(registry_number="77777-21")
        assert si_type.needs_review and si_type.created_by == self.verifier
        assert si_type.fif_raw["assistant_sources"][0].startswith("ФГИС")
        event = self.conversation.messages.filter(role="event").get()
        assert "заведён" in event.content and "проверки метролога" in event.content
        # повторно решить нельзя
        again = self.client_for(self.verifier).post(reverse("assistant_action_decide", args=[action.pk, "confirm"]))
        assert again.status_code == 400

    def test_reject(self):
        self.propose()
        action = PendingAction.objects.get()
        self.client_for(self.verifier).post(reverse("assistant_action_decide", args=[action.pk, "reject"]))
        action.refresh_from_db()
        assert action.status == "rejected" and not SiType.objects.filter(registry_number="77777-21").exists()

    def test_invalid_limits_are_rejected_before_card(self):
        tool_msg = self.propose(limits={**LIMITS, "q_min": 5})
        assert tool_msg.ui["error"] and not PendingAction.objects.exists()

    def test_sources_required(self):
        tool_msg = self.propose(sources=[])
        assert tool_msg.ui["error"] and "sources" in tool_msg.content

    def test_company_memory_needs_reviewer(self):
        fake = FakeLLM(call("memory_save", {"text": "Манометры — МПИ 1 год", "scope": "company"}), answer("ок"))
        self.send("запомни для всех", fake)
        action = PendingAction.objects.get()
        self.client_for(self.verifier).post(reverse("assistant_action_decide", args=[action.pk, "confirm"]))
        action.refresh_from_db()
        assert action.status == "failed" and not MemoryItem.objects.exists()

    def test_personal_memory_saved_directly_and_in_prompt(self):
        fake = FakeLLM(call("memory_save", {"text": "Мои объекты — Слободской район"}), answer("Запомнил"))
        self.send("запомни", fake)
        item = MemoryItem.objects.get()
        assert item.scope == Scope.PERSONAL and item.user == self.verifier
        prompt = build_system_prompt(self.verifier, tool_names=["mcp__rst__search_types"], mcp_errors=[])
        assert "Слободской район" in prompt
        other = build_system_prompt(self.metrologist, tool_names=[], mcp_errors=[])
        assert "Слободской район" not in other
        assert "RSTMetrology_MCP) сейчас недоступны" in other


# ---------------------------------------------------------------------------
class DocumentTests(Base):
    def test_protocol_requires_sources(self):
        fake = FakeLLM(call("protocol_draft", {
            "title": "ПРОТОКОЛ", "blocks": [{"type": "paragraph", "text": "x"}],
            "facts": [{"field": "Зав. №", "value": "190-157", "status": "confirmed"}],
        }), answer("ок"))
        self.send("протокол", fake)
        tool_msg = self.conversation.messages.get(role="tool")
        assert tool_msg.ui["error"] and "источник" in tool_msg.content

    @pytest.mark.skipif(shutil.which("typst") is None, reason="нет typst")
    def test_protocol_draft_makes_docx_and_pdf(self):
        fake = FakeLLM(call("protocol_draft", {
            "title": "ПРОТОКОЛ № ________", "subtitle": "первичной аттестации испытательного оборудования",
            "blocks": [
                {"type": "heading", "text": "1. Комиссия в составе:"},
                {"type": "fields", "rows": [{"label": "заводской номер", "value": "190-157"}]},
                {"type": "table", "columns": ["Точка", "Результат"], "rows": [["-70 °C", "________"]]},
                {"type": "signatures", "rows": [{"role": "Председатель", "name": "________"}]},
            ],
            "facts": [
                {"field": "Зав. №", "value": "190-157", "source": "со слов пользователя", "status": "confirmed"},
                {"field": "Результат -70", "value": "", "status": "missing"},
            ],
        }), answer("Черновик готов."))
        self.send("протокол", fake)
        tool_msg = self.conversation.messages.get(role="tool")
        assert not tool_msg.ui["error"], tool_msg.content
        names = [f["name"] for f in tool_msg.ui["files"]]
        assert names[0].endswith(".docx") and names[1].endswith(".pdf")
        assert tool_msg.ui["widgets"][0]["type"] == "facts"
        assert "Не заполнено полей: 1" in tool_msg.content
        pdf = Attachment.objects.get(name=names[1])
        assert pdf.file.read(4) == b"%PDF"

    def test_xlsx_and_docx_roundtrip(self):
        fake = FakeLLM(call("file_create", {
            "formats": ["xlsx", "docx"], "filename": "сводка",
            "sheets": [{"name": "Поверки", "columns": ["№", "Значение"], "rows": [[1, "2,5"]]}],
            "blocks": [{"type": "paragraph", "text": "Привет"}],
        }), answer("ок"))
        self.send("сделай файл", fake)
        files = Attachment.objects.filter(origin="agent").order_by("pk")
        assert [f.name for f in files] == ["сводка.xlsx", "сводка.docx"]
        docx_text = extract.extract_text(files[1].file.read(), "a.docx")
        assert "Привет" in docx_text
        xlsx_text = extract.extract_text(files[0].file.read(), "a.xlsx")
        assert "Поверки" in xlsx_text and "2.5" in xlsx_text

    def test_widget_validation(self):
        fake = FakeLLM(
            call("ui_show", {"widget": {"type": "bogus"}}, "a"),
            call("ui_show", {"widget": {"type": "table", "columns": ["a"], "rows": [["1"]]}}, "b"),
            answer("ок"),
        )
        self.send("покажи", fake)
        first, second = self.conversation.messages.filter(role="tool").order_by("pk")
        assert first.ui["error"] and not second.ui["error"]
        assert second.ui["widgets"][0]["type"] == "table"


# ---------------------------------------------------------------------------
class KnowledgeApiTests(Base):
    def test_verifier_cannot_upload_company_file_but_can_personal(self):
        client = self.client_for(self.verifier)
        upload = SimpleUploadedFile("правила.txt", "Пломбы ставим всегда".encode())
        response = client.post(reverse("assistant_knowledge"), {"file": upload, "scope": "company"})
        assert response.status_code == 403
        upload = SimpleUploadedFile("мои.txt", "Пломбы ставим всегда".encode())
        response = client.post(reverse("assistant_knowledge"), {"file": upload, "scope": "personal"})
        assert response.status_code == 201, response.data
        assert response.data["chunks"] == 1
        # личный файл не виден другим
        listing = self.client_for(self.metrologist).get(reverse("assistant_knowledge")).data
        assert listing["files"] == []

    def test_metrologist_uploads_docx_and_search_finds_it(self):
        import docx

        document = docx.Document()
        document.add_heading("Порядок оформления протоколов", 1)
        document.add_paragraph("Номер протокола выдаётся при подписании. Условия поверки берём из журнала погоды.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Эталон"
        table.rows[0].cells[1].text = "УПСЖ-3"
        buffer = io.BytesIO()
        document.save(buffer)
        client = self.client_for(self.metrologist)
        response = client.post(reverse("assistant_knowledge"), {
            "file": SimpleUploadedFile("порядок.docx", buffer.getvalue()), "category": "rules", "mode": "always",
        })
        assert response.status_code == 201, response.data
        kf = KnowledgeFile.objects.get()
        assert "| Эталон | УПСЖ-3 |" in kf.text
        hits = client.get(reverse("assistant_knowledge_search") + "?q=журнал погоды").data
        assert hits and hits[0]["file"] == "порядок"
        prompt = build_system_prompt(self.verifier, tool_names=[], mcp_errors=[])
        assert "УПСЖ-3" in prompt  # режим «всегда в контексте»
        client.patch(reverse("assistant_knowledge_detail", args=[kf.pk]), {"mode": "search"}, format="json")
        prompt = build_system_prompt(self.verifier, tool_names=[], mcp_errors=[])
        assert "УПСЖ-3" not in prompt and f"[{kf.pk}]" in prompt
        # поверитель общий файл не правит
        denied = self.client_for(self.verifier).patch(
            reverse("assistant_knowledge_detail", args=[kf.pk]), {"mode": "off"}, format="json")
        assert denied.status_code == 403

    def test_attachment_upload_extracts_text(self):
        response = self.client_for(self.verifier).post(reverse("assistant_attachment_upload"), {
            "file": SimpleUploadedFile("список.csv", "Зав. №;Тип\n123;СВК-15\n".encode("cp1251")),
        })
        assert response.status_code == 201
        assert Attachment.objects.get().text == "Зав. №\tТип\n123\tСВК-15"

    def test_old_doc_format_explained(self):
        response = self.client_for(self.verifier).post(reverse("assistant_attachment_upload"), {
            "file": SimpleUploadedFile("старый.doc", b"\xd0\xcf\x11\xe0"),
        })
        assert ".docx" in response.data["extract_error"]


# ---------------------------------------------------------------------------
class ChunkTests(TestCase):
    def test_chunks_cover_text_with_headings(self):
        text = "## Раздел 1\n" + ("Строка про манометры. " * 60) + "\n## Раздел 2\n" + ("Про весы. " * 300)
        chunks = extract.chunk_text(text)
        assert len(chunks) >= 3
        assert chunks[0][0] == "Раздел 1"
        assert any(h == "Раздел 2" for h, _ in chunks)
        assert all(len(body) <= extract.CHUNK_SIZE + 400 for _, body in chunks)


# ---------------------------------------------------------------------------
class FakeResponse:
    def __init__(self, body, *, content_type="application/json", status=200, headers=None):
        self.status_code = status
        self._body = body
        self.headers = {"Content-Type": content_type, **(headers or {})}
        self.text = body if isinstance(body, str) else json.dumps(body)

    def json(self):
        return self._body if not isinstance(self._body, str) else json.loads(self._body)

    def iter_lines(self, decode_unicode=False):
        for line in self.text.split("\n"):
            yield line if decode_unicode else line.encode()

    def close(self):
        pass


class McpClientTests(TestCase):
    def test_handshake_and_sse_tool_call(self):
        posts = []

        def fake_post(self_, url, headers=None, data=None, timeout=None, stream=None):
            payload = json.loads(data)
            posts.append((payload, dict(headers)))
            method = payload["method"]
            if method == "initialize":
                return FakeResponse({"jsonrpc": "2.0", "id": payload["id"], "result": {"protocolVersion": "2025-06-18"}},
                                    headers={"Mcp-Session-Id": "s-1"})
            if method == "notifications/initialized":
                return FakeResponse("", status=202)
            if method == "tools/call":
                event = json.dumps({"jsonrpc": "2.0", "id": payload["id"], "result": {
                    "content": [{"type": "text", "text": "mit_number: 16078-05"}], "isError": False}})
                return FakeResponse(f"event: message\ndata: {event}\n\n", content_type="text/event-stream")
            raise AssertionError(method)

        server = mock.Mock(url="http://rst:8080/mcp", headers={"Authorization": "Bearer t"})
        with mock.patch("requests.Session.post", fake_post), mock.patch("requests.Session.delete"):
            text, is_error = mcp_client.call(server, "search_types", {"search": "СГВ"})
        assert text == "mit_number: 16078-05" and not is_error
        assert posts[2][1]["Mcp-Session-Id"] == "s-1"
        assert posts[2][1]["Authorization"] == "Bearer t"
        assert mcp_client.tool_name("rst", "search_types") == "mcp__rst__search_types"

    @override_settings(RST_MCP_URL="http://rst:8080/mcp", RST_MCP_TOKEN="abc")
    def test_rst_server_seeded_from_env(self):
        from apps.assistant.models import McpServer
        from apps.assistant.setup import ensure_rst_server

        ensure_rst_server()
        server = McpServer.objects.get(name="rst")
        assert server.url == "http://rst:8080/mcp" and server.headers["Authorization"] == "Bearer abc"


class WebTests(TestCase):
    def test_parse_yandex_xml(self):
        xml = """<?xml version="1.0" encoding="utf-8"?><yandexsearch><response><results><grouping>
        <group><doc><url>https://fgis.gost.ru/x.pdf</url><title>Описание <hlword>типа</hlword></title>
        <passages><passage>Счётчики воды <hlword>СГВ</hlword></passage></passages></doc></group>
        </grouping></results></response></yandexsearch>""".encode()
        rows = web.parse_search_xml(xml)
        assert rows == [{"url": "https://fgis.gost.ru/x.pdf", "title": "Описание типа",
                         "snippet": "Счётчики воды СГВ", "date": ""}]

    def test_search_request(self):
        xml = b"<yandexsearch><response><results><grouping/></results></response></yandexsearch>"
        response = FakeResponse({"rawData": base64.b64encode(xml).decode()})
        with override_settings(YANDEX_VISION_API_KEY="k", YANDEX_VISION_FOLDER_ID="f"), \
                mock.patch("requests.post", return_value=response) as post:
            assert web.search("МИ 1592-2015 методика") == []
        body = post.call_args.kwargs["json"]
        assert body["folderId"] == "f" and body["query"]["queryText"] == "МИ 1592-2015 методика"

    def test_private_addresses_blocked(self):
        with pytest.raises(web.WebError):
            web.fetch("http://127.0.0.1:8000/admin/")
        with pytest.raises(web.WebError):
            web.fetch("file:///etc/passwd")


class LlmStreamTests(TestCase):
    @override_settings(OPENROUTER_API_KEY="k")
    def test_stream_assembles_text_and_tool_calls(self):
        events = [
            {"choices": [{"delta": {"content": "Ищу "}}]},
            {"choices": [{"delta": {"tool_calls": [{"index": 0, "id": "t1", "function": {"name": "web_", "arguments": "{\"qu"}}]}}]},
            {"choices": [{"delta": {"tool_calls": [{"index": 0, "function": {"name": "search", "arguments": "ery\": \"x\"}"}}]}}]},
            {"choices": [{"delta": {}, "finish_reason": "tool_calls"}], "usage": {"total_tokens": 7}},
        ]
        body = "\n".join(f"data: {json.dumps(e)}\n" for e in events) + "\ndata: [DONE]\n"
        streamed = []
        with mock.patch("requests.post", return_value=FakeResponse(body, content_type="text/event-stream")) as post:
            result = llm.chat("openrouter:a/b", [{"role": "user", "content": "x"}], on_text=streamed.append)
        assert post.call_args.kwargs["json"]["model"] == "a/b"
        assert result.text == "Ищу " and streamed == ["Ищу "]
        assert result.tool_calls[0].name == "web_search" and result.tool_calls[0].parsed() == {"query": "x"}
        assert result.usage == {"total_tokens": 7}

    @override_settings(YANDEX_VISION_API_KEY="k", YANDEX_VISION_FOLDER_ID="b1g")
    def test_yandex_model_uri(self):
        with mock.patch("requests.post", return_value=FakeResponse("data: [DONE]\n", content_type="text/event-stream")) as post:
            llm.chat("yandex:qwen3-235b-a22b-fp8/latest", [{"role": "user", "content": "x"}])
        assert post.call_args.kwargs["json"]["model"] == "gpt://b1g/qwen3-235b-a22b-fp8/latest"
        assert post.call_args.kwargs["headers"]["Authorization"] == "Api-Key k"
